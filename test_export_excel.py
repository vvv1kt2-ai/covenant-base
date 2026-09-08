"""Tests for export_excel — the Excel deliverable contract.

Covers: category column reads the model's unified taxonomy verbatim,
error rows keep the parsing-failure marker, zero-covenant rows keep the
"put absent" row type, program-sourced rows keep the green tint.
"""
import shutil
import uuid
from pathlib import Path

import pytest
from openpyxl import load_workbook

from covenant_models import Covenant, ResultEntry, save_results
from export_excel import export_to_excel

HEADERS = [
    "Эмитент", "Выпуск", "ISIN", "Рейтинг",
    "Ковенантов\nв выпуске",
    "№", "Категория ковенанта", "Суть ковенанта",
    "Документ", "Пункт", "Стр.", "Файл",
    "Цитата из эмиссионной документации",
    "Источник\n(Финам)",
]
CATEGORY_COL = 7
ISIN_COL = 3
DOCUMENT_COL = 9


@pytest.fixture
def work_dir():
    """Temp dir for fixtures and the exported workbook, self-cleaning."""
    d = Path(f"covtest-{uuid.uuid4().hex[:8]}")
    d.mkdir()
    yield d
    shutil.rmtree(d, ignore_errors=True)


def export_entries(entries, work_dir):
    """Write entries via the model's I/O, then export to Excel."""
    src = work_dir / "results.json"
    save_results(src, entries)
    out = work_dir / "out.xlsx"
    export_to_excel(results_path=str(src), output_path=str(out))
    return load_workbook(out).active


def make_entries():
    with_cov = ResultEntry(
        isin="RU000A1089A3",
        issuer="АО «Эмитент»",
        issue_name="ЭМИТЕНТ-БО-01",
        rating="A-(RU)",
        decision_url="https://bonds.finam.ru/issue/detailsABC/default.asp",
        processed_at="2026-01-01T00:00:00",
    )
    with_cov.add_covenant(Covenant(
        category="Делистинг",
        essence="В случае делистинга облигаций",
        document="Решение о выпуске",
        section="п. 5.6.1, 1",
        page=12,
        quote="Владельцы имеют право требовать досрочного погашения",
    ))

    with_err = ResultEntry(
        isin="RU000A10ERR1",
        parse_errors=["Failed to download decision PDF"],
        processed_at="2026-01-01T00:00:00",
    )

    zero = ResultEntry(
        isin="RU000A10ZERO",
        issuer="АО «Пустой»",
        decision_url="https://bonds.finam.ru/issue/detailsXYZ/default.asp",
        processed_at="2026-01-01T00:00:00",
    )
    return [with_cov, with_err, zero]


def test_header_contract(work_dir):
    ws = export_entries(make_entries(), work_dir)
    for col, text in enumerate(HEADERS, 1):
        assert ws.cell(row=1, column=col).value == text


def test_category_column_reads_model_taxonomy(work_dir):
    ws = export_entries(make_entries(), work_dir)

    # Row 2: covenant row — category verbatim from the model
    assert ws.cell(row=2, column=ISIN_COL).value == "RU000A1089A3"
    assert ws.cell(row=2, column=CATEGORY_COL).value == "Делистинг"

    # Row 3: error ISIN — parsing-failure marker
    assert ws.cell(row=3, column=ISIN_COL).value == "RU000A10ERR1"
    assert ws.cell(row=3, column=CATEGORY_COL).value == "⚠ ОШИБКА ПАРСИНГА"

    # Row 4: zero-covenant ISIN — put-absent row type
    assert ws.cell(row=4, column=ISIN_COL).value == "RU000A10ZERO"
    assert ws.cell(row=4, column=CATEGORY_COL).value == "Положения о досрочном погашении"


def test_program_covenant_category_and_tint(work_dir):
    entries = make_entries()
    entries[0].add_covenant(Covenant(
        category="Кросс-дефолт",
        essence="Кросс-дефолт по программе",
        document="Программа облигаций (e-disclosure)",
    ))

    ws = export_entries(entries, work_dir)

    # Covenant #1 (decision) at Excel row 2, covenant #2 (program) at row 3
    assert ws.cell(row=2, column=CATEGORY_COL).value == "Делистинг"
    assert ws.cell(row=3, column=CATEGORY_COL).value == "Кросс-дефолт"
    assert ws.cell(row=3, column=DOCUMENT_COL).value == "Программа облигаций (e-disclosure)"

    # Green tint on the program-sourced row, none on the decision row
    assert ws.cell(row=3, column=CATEGORY_COL).fill.start_color.rgb in ("00E2EFDA", "FFE2EFDA")
    assert ws.cell(row=2, column=CATEGORY_COL).fill.fill_type != "solid"
