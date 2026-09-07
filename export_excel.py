"""Export results.json to Excel following the template format."""
import json
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: pip install openpyxl")
    exit(1)


def export_to_excel(results_path: str = "results.json", output_path: str = "covarianants.xlsx"):
    with open(results_path, encoding="utf-8") as f:
        data = json.load(f)

    wb = Workbook()
    ws = wb.active
    ws.title = "Ковенанты"

    # --- Styles ---
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    wrap = Alignment(wrap_text=True, vertical="top")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    link_font = Font(color="0563C1", underline="single")
    warn_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    warn_font = Font(bold=True, color="C00000")
    # Light green tint for program-sourced covenants
    prog_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

    # --- Headers ---
    headers = [
        "Эмитент", "Выпуск", "ISIN", "Рейтинг",
        "Ковенантов\nв выпуске",
        "№", "Категория ковенанта", "Суть ковенанта",
        "Документ", "Пункт", "Стр.", "Файл",
        "Цитата из эмиссионной документации",
        "Источник\n(Финам)",
    ]
    for col, text in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap
        cell.border = thin_border

    # --- Column widths ---
    widths = [25, 20, 16, 8, 12, 5, 30, 45, 18, 14, 6, 35, 60, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # --- Data ---
    row = 2
    total_with_covenants = 0
    total_without = 0
    total_errors = 0

    for r in data:
        issuer = r.get("issuer", "")
        issue_name = r.get("issue_name", "")
        isin = r.get("isin", "")
        rating = r.get("rating", "")
        covenants = r.get("covenants", [])
        num_covenants = r.get("total_covenants", len(covenants))
        decision_url = r.get("decision_url", "")
        filename = decision_url.split("/")[-1] if decision_url else ""
        program_url = r.get("program_url", "")
        parse_errors = r.get("parse_errors", [])

        # --- Error ISINs: mark for manual check ---
        has_manual_check = r.get("requires_manual_check", False)
        error_reason = r.get("manual_check_reason", "")
        if (parse_errors and num_covenants == 0) or has_manual_check:
            total_errors += 1
            if has_manual_check:
                error_msg = error_reason or "Программа облигаций скачана, требует ручной проверки"
            else:
                error_msg = "; ".join(parse_errors)[:300]

            values = [
                issuer, issue_name, isin, rating,
                "Требует\nручной проверки",
                "—",
                "⚠ ОШИБКА ПАРСИНГА",
                f"Не удалось извлечь ковенанты: {error_msg}",
                "Решение о выпуске",
                "",
                "",
                filename,
                "Требует ручной проверки",
                decision_url,
            ]

            for col_idx, value in enumerate(values, 1):
                cell = ws.cell(row=row, column=col_idx, value=value)
                cell.alignment = wrap
                cell.border = thin_border
                cell.fill = warn_fill
                if col_idx in (6, 7, 13):
                    cell.font = warn_font

            if decision_url:
                link_cell = ws.cell(row=row, column=14)
                link_cell.hyperlink = decision_url
                link_cell.value = decision_url
                link_cell.font = link_font

            row += 1
            continue

        # --- ISINs with covenants ---
        if num_covenants > 0:
            total_with_covenants += 1

            for i, cov in enumerate(covenants):
                is_provided = cov.get("is_provided", False)
                is_from_program = "Программа" in cov.get("document", "")

                if is_provided:
                    category = "Информационный (отчётность эмитента) +\nРаскрытие отчётности эмитента"
                else:
                    category = "Положения о досрочном погашении"

                # Essence comes directly from parser (event title or section title)
                essence = cov.get("essence", "")
                section = cov.get("section", "")
                page = cov.get("page", "")
                quote = cov.get("quote", "").replace("\n", " ")
                if len(quote) > 500:
                    quote = quote[:500] + "\u2026"

                # File: show program filename for program-sourced covenants
                if is_from_program and program_url:
                    file_name = program_url.split("/")[-1].split("?")[0]
                else:
                    file_name = filename

                # Source URL: program or decision
                source_url = program_url if (is_from_program and program_url) else decision_url

                values = [
                    issuer if i == 0 else "",
                    issue_name if i == 0 else "",
                    isin if i == 0 else "",
                    rating if i == 0 else "",
                    num_covenants if i == 0 else "",
                    cov.get("number", i + 1),
                    category,
                    essence,
                    cov.get("document", ""),
                    section,
                    page,
                    file_name,
                    quote,
                    source_url,
                ]

                for col_idx, value in enumerate(values, 1):
                    cell = ws.cell(row=row, column=col_idx, value=value)
                    cell.alignment = wrap
                    cell.border = thin_border
                    # Green tint for program-sourced covenants
                    if is_from_program:
                        cell.fill = prog_fill

                # Hyperlink to source
                if source_url:
                    link_cell = ws.cell(row=row, column=14)
                    link_cell.hyperlink = source_url
                    link_cell.value = source_url
                    link_cell.font = link_font

                row += 1
        else:
            # 0 covenants, no errors — put option absent
            total_without += 1

            values = [
                issuer, issue_name, isin, rating,
                0,
                "\u2014",
                "Положения о досрочном погашении",
                "Put-опция у владельцев отсутствует",
                "Решение о выпуске",
                "",
                "",
                filename,
                "",
                decision_url,
            ]

            for col_idx, value in enumerate(values, 1):
                cell = ws.cell(row=row, column=col_idx, value=value)
                cell.alignment = wrap
                cell.border = thin_border

            if decision_url:
                link_cell = ws.cell(row=row, column=14)
                link_cell.hyperlink = decision_url
                link_cell.value = decision_url
                link_cell.font = link_font

            row += 1

    # --- Summary ---
    data_end_row = row - 1
    row += 1
    ws.cell(row=row, column=1, value="ИТОГ").font = header_font
    ws.cell(row=row, column=2, value=f"ISIN: {len(data)}")
    row += 1
    ws.cell(row=row, column=1, value="С ковенантами:").font = header_font
    ws.cell(row=row, column=2, value=total_with_covenants)
    row += 1
    ws.cell(row=row, column=1, value="Без ковенантов:").font = header_font
    ws.cell(row=row, column=2, value=total_without)
    row += 1
    ws.cell(row=row, column=1, value="Требует ручной проверки:").font = header_font
    ws.cell(row=row, column=2, value=total_errors)

    ws.auto_filter.ref = f"A1:N{data_end_row}"

    wb.save(output_path)
    print(f"Saved: {output_path}")
    print(f"  Rows: {data_end_row - 1} (excl. header)")
    print(f"  ISINs with covenants: {total_with_covenants}")
    print(f"  ISINs without covenants: {total_without}")
    print(f"  ISINs with errors (manual check): {total_errors}")


if __name__ == "__main__":
    import sys as _sys
    out = _sys.argv[1] if len(_sys.argv) > 1 else "covarianants.xlsx"
    export_to_excel(output_path=out)